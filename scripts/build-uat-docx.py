from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/phase-11-12/BB_UAT_ky-xac-nhan.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "64748B"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
AMBER_FILL = "FFF7E0"
RED_FILL = "FDECEC"
GREEN_FILL = "EAF7EF"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def font_run(run, size=11, color="000000", bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.append(begin)
    run.append(instr)
    run.append(separate)
    run.append(text)
    run.append(end)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font_run(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font_run(p.add_run(text))
    return p


def add_callout(doc, title, text, fill=LIGHT_FILL, color=INK):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    font_run(p.add_run(title + "  "), size=10.5, color=color, bold=True)
    font_run(p.add_run(text), size=10.5, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        set_cell_shading(header.cells[index], header_fill)
        p = header.cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        font_run(p.add_run(value), size=font_size, color=INK, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if index == 0 and len(widths) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font_run(p.add_run(str(value)), size=font_size, color="1F2937")
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
font_run(hp.add_run("ĐIỆN LẠNH 247  ·  UAT & GO-LIVE"), size=8.5, color=MUTED, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font_run(fp.add_run("BB_UAT_ky-xac-nhan  ·  Trang "), size=8.5, color=MUTED)
add_field(fp, "PAGE")
font_run(fp.add_run(" / "), size=8.5, color=MUTED)
add_field(fp, "NUMPAGES")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(8)
font_run(p.add_run("BIÊN BẢN NGHIỆM THU NGƯỜI DÙNG"), size=10, color=BLUE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
font_run(p.add_run("BB_UAT_ky-xac-nhan"), size=25, color=INK, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
font_run(p.add_run("Điện Lạnh 247 · Service-only platform · Release v1.0.0"), size=13, color=MUTED)

metadata = [
    ("Phạm vi", "Giai đoạn 11 — UAT, sửa lỗi; Giai đoạn 12 — go-live và bàn giao"),
    ("Môi trường", "Staging gần production (điền URL và release SHA khi thực hiện)"),
    ("Ngày UAT", "____ / ____ / ______    Khung giờ: __________"),
    ("Điều phối", "____________________________________________"),
    ("Quyết định", "NO-GO cho đến khi đủ chữ ký và P0/P1 = 0"),
]
table = doc.add_table(rows=0, cols=2)
table.style = "Table Grid"
for label, value in metadata:
    cells = table.add_row().cells
    set_cell_shading(cells[0], HEADER_FILL)
    p0, p1 = cells[0].paragraphs[0], cells[1].paragraphs[0]
    p0.paragraph_format.space_after = p1.paragraph_format.space_after = Pt(0)
    font_run(p0.add_run(label), size=10, color=INK, bold=True)
    font_run(p1.add_run(value), size=10, color="1F2937")
set_table_geometry(table, [1900, 7460])
set_repeat_table_header(table.rows[0])

doc.add_paragraph()
add_callout(
    doc,
    "TRẠNG THÁI KIỂM SOÁT",
    "Tài liệu này là mẫu ký thật. CI PASS không thay thế UAT không hướng dẫn, staging drill hoặc chữ ký của người tiếp nhận.",
    fill=AMBER_FILL,
    color="7A5A00",
)

doc.add_heading("1. Nguyên tắc nghiệm thu", level=1)
for item in (
    "Người tham gia chỉ nhận mục tiêu nghiệp vụ, không được chỉ vị trí từng nút.",
    "Điều phối viên chỉ quan sát; chỉ can thiệp khi có nguy cơ mất dữ liệu hoặc an toàn.",
    "Không dùng dữ liệu production. Ảnh, số điện thoại, địa chỉ và tài khoản đều là dữ liệu staging tổng hợp.",
    "P0/P1 phải bằng 0 trước quyết định GO; mọi lỗi phải có mã, bằng chứng đã redact và người chịu trách nhiệm.",
):
    add_bullet(doc, item)

doc.add_heading("2. Người tham gia và phạm vi quyền", level=1)
add_table(
    doc,
    ["Vai trò", "Người thực hiện", "Tài khoản staging", "Xác nhận không dùng production"],
    [
        ("Khách hàng", "________________", "________________", "PASS / FAIL"),
        ("Nhân viên vận hành", "________________", "________________", "PASS / FAIL"),
        ("Người quản trị", "________________", "________________", "PASS / FAIL"),
        ("QA quan sát", "________________", "Không đăng nhập thay", "PASS / FAIL"),
    ],
    [1700, 2300, 2900, 2460],
)

doc.add_heading("3. UAT theo mục tiêu — không hướng dẫn từng nút", level=1)
doc.add_heading("3.1 Khách hàng", level=2)
add_callout(doc, "MỤC TIÊU", "Đặt lịch sửa điều hòa, tải ảnh hiện trạng, tra cứu mã yêu cầu và đổi lịch phù hợp.", fill=LIGHT_FILL)
add_table(
    doc,
    ["#", "Quan sát bắt buộc", "Kết quả", "Ghi chú / mã lỗi"],
    [
        ("1", "Tự tìm được CTA đặt lịch; hiểu quy trình 4 bước.", "PASS / FAIL", "________________"),
        ("2", "Hiểu giá chỉ tham khảo và lịch chỉ chốt sau xác nhận.", "PASS / FAIL", "________________"),
        ("3", "Upload ảnh hợp lệ; lỗi ảnh không làm mất yêu cầu.", "PASS / FAIL", "________________"),
        ("4", "Nhấn gửi lặp không tạo yêu cầu thứ hai.", "PASS / FAIL", "________________"),
        ("5", "Tự tra cứu, xem timeline và đổi lịch đúng điều kiện.", "PASS / FAIL", "________________"),
    ],
    [520, 4860, 1420, 2560],
)

doc.add_heading("3.2 Nhân viên vận hành", level=2)
add_callout(doc, "MỤC TIÊU", "Tìm yêu cầu mới, xác nhận, chọn kỹ thuật viên đúng kỹ năng/khu vực và xử lý xung đột lịch.", fill=LIGHT_FILL)
add_table(
    doc,
    ["#", "Quan sát bắt buộc", "Kết quả", "Ghi chú / mã lỗi"],
    [
        ("1", "Navigation chỉ hiện nghiệp vụ đúng quyền.", "PASS / FAIL", "________________"),
        ("2", "Xác nhận và phân công không cần quyền Settings/Audit.", "PASS / FAIL", "________________"),
        ("3", "Backend chặn hai lịch trùng bằng 409.", "PASS / FAIL", "________________"),
        ("4", "Mất mạng: dừng thao tác, banner rõ ràng, kết nối lại an toàn.", "PASS / FAIL", "________________"),
        ("5", "Thông báo provider lỗi không làm mất nghiệp vụ chính.", "PASS / FAIL", "________________"),
    ],
    [520, 4860, 1420, 2560],
)

doc.add_heading("3.3 Người quản trị", level=2)
add_callout(doc, "MỤC TIÊU", "Kiểm tra SLA, báo giá, CMS/ảnh, audit, cấu hình bảo mật và phục hồi dữ liệu.", fill=LIGHT_FILL)
add_table(
    doc,
    ["#", "Quan sát bắt buộc", "Kết quả", "Ghi chú / mã lỗi"],
    [
        ("1", "Step-up bắt buộc trước thay đổi nhạy cảm.", "PASS / FAIL", "________________"),
        ("2", "Audit không chứa token, cookie, địa chỉ hoặc ảnh nhạy cảm.", "PASS / FAIL", "________________"),
        ("3", "CMS upload sai magic-byte bị từ chối trước storage.", "PASS / FAIL", "________________"),
        ("4", "Tự đọc health/log/alert và tìm request ID.", "PASS / FAIL", "________________"),
        ("5", "Tự chạy encrypted backup/restore và rollback image.", "PASS / FAIL", "________________"),
    ],
    [520, 4860, 1420, 2560],
)

doc.add_heading("4. Ma trận lỗi bắt buộc", level=1)
add_table(
    doc,
    ["Tình huống", "Kỳ vọng nghiệm thu", "Bằng chứng", "Kết quả"],
    [
        ("API timeout", "Thông báo dễ hiểu; retry giữ idempotency key.", "Trace / screenshot", "PASS / FAIL"),
        ("Mất mạng", "Giữ dữ liệu form; banner offline/restored; mutation không retry mù.", "Playwright", "PASS / FAIL"),
        ("Upload lỗi", "Sai MIME/magic/size bị chặn; không lộ stack.", "400 + unit", "PASS / FAIL"),
        ("Submit nhiều lần", "Cùng khóa trả cùng mã; payload khác 409.", "API log", "PASS / FAIL"),
        ("Token hết hạn", "401; refresh rotation; revoked session không dùng lại.", "API log", "PASS / FAIL"),
        ("Xung đột lịch", "Backend trả 409 cho lịch trùng.", "API log", "PASS / FAIL"),
        ("Provider lỗi", "Business commit; outbox FAILED rồi DEAD.", "Unit / DB", "PASS / FAIL"),
    ],
    [1850, 3970, 1800, 1740],
    font_size=9,
)

doc.add_heading("5. Sổ lỗi P0/P1", level=1)
add_table(
    doc,
    ["Mã", "Mức", "Mô tả / bước tái hiện", "Owner", "Trạng thái", "Bằng chứng retest"],
    [("____", "P0 / P1", "________________________________", "________", "OPEN / FIXED", "________________") for _ in range(6)],
    [700, 900, 3200, 1200, 1500, 1860],
    font_size=8.7,
)

add_callout(doc, "ĐIỀU KIỆN CHẶN", "Nếu còn bất kỳ P0 hoặc P1 ở trạng thái OPEN, quyết định mặc định là NO-GO.", fill=RED_FILL, color="9B1C1C")

doc.add_heading("6. Bằng chứng kỹ thuật và staging drill", level=1)
add_table(
    doc,
    ["Cổng", "Run ID / artifact / đường dẫn", "Kết quả", "Người xác minh"],
    [
        ("Lint + typecheck + unit", "____________________________", "PASS / FAIL", "____________"),
        ("Integration + architecture + Mock", "____________________________", "PASS / FAIL", "____________"),
        ("Playwright + production build", "____________________________", "PASS / FAIL", "____________"),
        ("Security/secret + RBAC negative", "____________________________", "PASS / FAIL", "____________"),
        ("Encrypted backup/restore staging", "Checksum: __________________", "PASS / FAIL", "____________"),
        ("RPO / RTO", "RPO: ______  RTO: ______", "PASS / FAIL", "____________"),
    ],
    [2450, 3500, 1500, 1910],
)

doc.add_heading("7. Go-live readiness", level=1)
for item in (
    "Domain và DNS đã chốt; certificate CA bao phủ cả customer/admin domain.",
    "Secrets, SMTP, storage và production database được inject qua kênh an toàn.",
    "SERVICE_ONLY_MODE=true; RUN_SEED=false; demo/dev/mock đều tắt.",
    "Health monitor, alert escalation, backup schedule và off-host retention hoạt động.",
    "Tài khoản bootstrap đã tắt; người nhận đổi mật khẩu, bật 2FA và revoke quyền thừa.",
    "Rollback image, migration compatibility và restore approval hai người đã diễn tập.",
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    font_run(p.add_run("☐  "), size=12, color=BLUE, bold=True)
    font_run(p.add_run(item), size=10.5)

doc.add_heading("8. Bàn giao tự vận hành", level=1)
add_table(
    doc,
    ["Bài thực hành", "Người nhận tự hoàn tất", "Người quan sát", "Kết quả"],
    [
        ("Quản lý yêu cầu và phân công", "________________", "________________", "PASS / FAIL"),
        ("Báo giá, nghiệm thu, bảo hành", "________________", "________________", "PASS / FAIL"),
        ("CMS và thay ảnh theo asset key", "________________", "________________", "PASS / FAIL"),
        ("Audit, health, log và alert", "________________", "________________", "PASS / FAIL"),
        ("Backup, restore và rollback", "________________", "________________", "PASS / FAIL"),
    ],
    [3100, 2200, 2200, 1860],
)

doc.add_heading("9. Quyết định và chữ ký", level=1)
add_callout(doc, "QUYẾT ĐỊNH", "☐ GO    ☐ NO-GO    ☐ GO CÓ ĐIỀU KIỆN (không áp dụng nếu còn P0/P1)", fill=GREEN_FILL, color="1F3A5F")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
font_run(p.add_run("Điều kiện / ghi chú quyết định: "), bold=True, color=INK)
font_run(p.add_run("________________________________________________________________________________\n________________________________________________________________________________"))

add_table(
    doc,
    ["Vai trò ký", "Họ tên", "Ngày giờ", "Chữ ký / xác nhận"],
    [
        ("Đại diện chủ sở hữu", "________________", "________________", "________________"),
        ("Đại diện vận hành", "________________", "________________", "________________"),
        ("Đại diện QA/UAT", "________________", "________________", "________________"),
        ("Người tiếp nhận", "________________", "________________", "________________"),
    ],
    [2350, 2400, 2100, 2510],
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font_run(p.add_run("Tài liệu chỉ có hiệu lực bàn giao khi đủ bốn chữ ký và release gate trả PASS."), size=9.5, color=MUTED, italic=True)

doc.core_properties.title = "BB UAT ký xác nhận — Điện Lạnh 247 v1.0.0"
doc.core_properties.subject = "Giai đoạn 11–12: UAT, go-live và bàn giao"
doc.core_properties.author = "Điện Lạnh 247"
doc.core_properties.keywords = "UAT, go-live, handover, v1.0.0"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
